"""Independent verification from original bytes. Deliberately never imports core.

The oracle uses whitespace chunks containing an Arabic ordinary letter, justified
by the inventory of this input (editorial signs are standalone chunks). This is
independent of the production tokenizer, rather than a second invocation of it.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import csv
import gzip
import hashlib
import json
import math
import re
import sqlite3
import unicodedata as ud

SOURCE_SHA256 = "295faa755008d0a292d2e285368a06dc478bf37cb5f29ed29bb06d7c6ca0bf0a"
PROFILES = ("raw", "nfc", "plain", "search")
REMOVED = set(range(0x064B, 0x0653)) | {0x0670, 0x0640} | set(range(0x06D6, 0x06DD)) | {0x06DE, 0x06E9}
SEARCH_TRANSLATION = str.maketrans({c: "ا" for c in "أإآٱ"} | {"ى": "ي"})


def independent_profiles(text: str) -> dict:
    nfc = ud.normalize("NFC", text)
    plain = "".join(c for c in nfc if ord(c) not in REMOVED)
    return {"raw": text, "nfc": nfc, "plain": plain, "search": plain.translate(SEARCH_TRANSLATION)}


def independent_source(text: str) -> dict:
    """Read raw line layout and whitespace chunks without production code."""
    verses, tokens, issues = [], [], []
    offset = 0
    surah_positions = Counter()
    for line_number, line in enumerate(text.splitlines(keepends=True), 1):
        body = line.rstrip("\r\n")
        parts = body.split("|", 2)
        if len(parts) != 3 or not parts[0].isdigit() or not parts[1].isdigit():
            issues.append({"line": line_number, "issue": "invalid_record"})
            offset += len(line)
            continue
        surah, ayah = map(int, parts[:2])
        raw = parts[2]
        begin = offset + len(parts[0]) + len(parts[1]) + 2
        verse_id = f"{surah}:{ayah}"
        candidates = []
        # split() supplies the actual oracle units; find() only recovers coordinates.
        cursor = 0
        for chunk in raw.split():
            local_start = raw.find(chunk, cursor)
            cursor = local_start + len(chunk)
            if not any(ud.category(c) == "Lo" for c in chunk):
                continue
            candidates.append((chunk, begin + local_start, begin + cursor))
            if any(ord(c) >= 0x06D6 for c in chunk):
                issues.append({"verse_id": verse_id, "issue": "editorial_mark_in_lexical_chunk"})
        leading = [independent_profiles(x[0])["plain"] for x in candidates[:4]]
        prefix = ayah == 1 and surah not in (1, 9) and leading == ["بسم", "الله", "الرحمن", "الرحيم"]
        verse_tokens = []
        for pos, (chunk, start, end) in enumerate(candidates, 1):
            surah_positions[surah] += 1
            token = {
                "token_id": f"{verse_id}:{pos}", "verse_id": verse_id,
                "surah_id": surah, "ayah_id": ayah,
                "global_ayah": len(verses) + 1, "global_token": len(tokens) + 1,
                "token_in_ayah": pos, "token_in_surah": surah_positions[surah],
                **independent_profiles(chunk), "raw_start": start, "raw_end": end,
                "letter_count": sum(ud.category(c) == "Lo" for c in chunk),
                "mark_count": sum(ud.category(c).startswith("M") for c in chunk),
                "is_opening_basmala": int(prefix and pos <= 4),
            }
            tokens.append(token)
            verse_tokens.append(token)
        verses.append({
            "verse_id": verse_id, "surah_id": surah, "ayah_id": ayah,
            "global_ayah": len(verses) + 1, "raw_text": raw,
            **{k: v for k, v in independent_profiles(raw).items() if k != "raw"},
            "raw_start": begin, "raw_end": begin + len(raw),
            "token_count": len(verse_tokens),
            "letter_count": sum(t["letter_count"] for t in verse_tokens),
            "mark_count_all": sum(ud.category(c).startswith("M") for c in raw),
            "mark_count_tokens": sum(t["mark_count"] for t in verse_tokens),
        })
        offset += len(line)
    return {"verses": verses, "tokens": tokens, "issues": issues}


def _record(checks: list, check_id: str, condition: bool | None, method: str, **details):
    checks.append({"check_id": check_id, "status": "not_run" if condition is None else "pass" if condition else "fail", "method_ru": method, **details})


def _rows(path: Path):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        yield from csv.DictReader(handle)


def validate_database(db_path: Path, source_text: str, reference: dict, checks: list) -> dict:
    if not db_path.exists():
        _record(checks, "database.present", False, "Проверка существования основной SQLite БД.")
        return {}
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
    _record(checks, "database.integrity", integrity == "ok", "Встроенная проверка целостности SQLite.", observed=integrity)
    found = {}
    for table, key in (("verses", "verse_id"), ("tokens", "token_id")):
        stored = [dict(row) for row in conn.execute(f"SELECT * FROM {table} ORDER BY " + ("global_ayah" if table == "verses" else "global_token"))]
        expected = reference[table]
        found[table] = stored
        _record(checks, f"{table}.row_count", len(stored) == len(expected), "Независимый разбор исходных строк и лексических пробельных фрагментов.", observed=len(stored), expected=len(expected))
        _record(checks, f"{table}.unique_ids", len({r[key] for r in stored}) == len(stored), "Уникальность идентификатора; одинаковый текст разрешён.")
        mismatches = []
        spans = []
        missing_fields = set(expected[0] if expected else ()) - set(stored[0] if stored else ())
        missing_fields -= {"mark_count_all", "mark_count_tokens"}
        for i, (row, ref) in enumerate(zip(stored, expected)):
            raw_field = "raw_text" if table == "verses" else "raw"
            if source_text[row["raw_start"]:row["raw_end"]] != row[raw_field]:
                spans.append(row[key])
            # A verse's mark_count is audited separately: editorial marks are
            # outside tokens, so the exact convention must not be guessed here.
            for field, value in ref.items():
                if field in row and row[field] != value:
                    if len(mismatches) < 20:
                        mismatches.append({"id": row[key], "field": field, "stored": row[field], "expected": value})
        _record(checks, f"{table}.raw_spans", not spans and len(stored) == len(expected), "Восстановление каждого raw-фрагмента по Unicode-смещению в неизменной строке файла.", mismatch_ids=spans[:20])
        _record(checks, f"{table}.independent_values", not mismatches and not missing_fields and len(stored) == len(expected), "Сопоставление всех адресов, последовательных индексов, профилей, длин и басмальных флагов с независимым разбором.", mismatches=mismatches, missing_fields=sorted(missing_fields))
    tok = found["tokens"]
    verses = found["verses"]
    agg = Counter()
    for row in tok:
        agg[(row["verse_id"], "token_count")] += 1
        for field in ("letter_count", "mark_count"):
            agg[(row["verse_id"], field)] += row[field]
    token_bad = [v["verse_id"] for v in verses if v["token_count"] != agg[(v["verse_id"], "token_count")] or v["letter_count"] != agg[(v["verse_id"], "letter_count")]]
    _record(checks, "aggregates.verse_token_letter", not token_bad, "Суммы числа токенов и базовых букв по токенам равны значениям аятов.", mismatch_ids=token_bad[:20])
    mark_modes = set()
    bad_marks = []
    for v, ref in zip(verses, reference["verses"]):
        if v["mark_count"] == ref["mark_count_all"]:
            mark_modes.add("all_verse_marks")
        else:
            bad_marks.append(v["verse_id"])
    _record(checks, "aggregates.mark_conservation", not bad_marks, "Диакритики аятов сверены отдельно от токенов: самостоятельные редакторские знаки не входят в лексические токены.", conventions_observed=sorted(mark_modes), mismatch_ids=bad_marks[:20], source_all_marks=sum(v["mark_count_all"] for v in reference["verses"]), source_token_marks=sum(t["mark_count"] for t in reference["tokens"]))
    profile_counts = {}
    for scope in ("file", "numbered"):
        selected = [t for t in tok if scope == "file" or not t["is_opening_basmala"]]
        independent = [t for t in reference["tokens"] if scope == "file" or not t["is_opening_basmala"]]
        for profile in PROFILES:
            expected_freq = Counter(t[profile] for t in independent)
            sql = f"SELECT {profile}, COUNT(*) FROM tokens" + (" WHERE is_opening_basmala = 0" if scope == "numbered" else "") + f" GROUP BY {profile}"
            actual = dict(conn.execute(sql))
            _record(checks, f"frequencies.{scope}.{profile}", actual == dict(expected_freq), "Полный частотный словарь SQL сверяется с независимым Counter по исходным вхождениям.", total=sum(actual.values()), types=len(actual))
            profile_counts[f"{scope}.{profile}"] = {"tokens": len(selected), "types": len(actual)}
    _validate_relational_indexes(conn, source_text, reference, checks)
    conn.close()
    return {"profile_counts": profile_counts, "tables": found}


def _validate_relational_indexes(conn, source_text, reference, checks):
    """Validate the stored indexes themselves, not only re-queried token totals."""
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    required = {"spans", "source_records", "surahs", "token_scopes", "vocabulary", "occurrences"}
    missing = required - tables
    _record(checks, "database.complete_schema", not missing if not missing else None,
            "Наличие вспомогательных таблиц сохранения исходника и полных индексов.", missing_tables=sorted(missing))
    foreign = [tuple(row) for row in conn.execute("PRAGMA foreign_key_check")]
    _record(checks, "database.foreign_keys", not foreign, "SQLite foreign_key_check всех таблиц.", violations=foreign[:20])
    if "spans" in tables:
        by_verse = defaultdict(list)
        for row in conn.execute("SELECT * FROM spans ORDER BY raw_start"):
            by_verse[row["verse_id"]].append(dict(row))
        errors = []
        for verse in reference["verses"]:
            spans = by_verse.pop(verse["verse_id"], [])
            cursor = verse["raw_start"]
            good = True
            for span in spans:
                good &= span["raw_start"] == cursor and span["raw_end"] > cursor
                good &= source_text[span["raw_start"]:span["raw_end"]] == span["raw"]
                cursor = span["raw_end"]
            good &= cursor == verse["raw_end"] and "".join(s["raw"] for s in spans) == verse["raw_text"]
            if not good:
                errors.append(verse["verse_id"])
        _record(checks, "source.spans_partition", not errors and not by_verse,
                "Все spans без разрывов и перекрытий восстанавливают весь raw-текст каждого аята.", mismatches=errors[:20], unexpected_verses=list(by_verse)[:20])
    if "source_records" in tables:
        rows = [dict(r) for r in conn.execute("SELECT * FROM source_records ORDER BY source_line")]
        lines = source_text.splitlines(keepends=True)
        cursor, errors = 0, []
        for i, (row, line) in enumerate(zip(rows, lines), 1):
            expected_id = ":".join(line.split("|", 2)[:2])
            ending = line[len(line.rstrip("\r\n")):]
            if row != {"verse_id": expected_id, "source_line": i, "line_start": cursor, "line_end": cursor + len(line), "line_ending": ending}:
                errors.append(i)
            cursor += len(line)
        _record(checks, "source.records_partition", len(rows) == len(lines) and cursor == len(source_text) and not errors,
                "Координаты всех исходных записей включают метаданные и точные переводы строк.", mismatch_lines=errors[:20])
    if "surahs" in tables:
        rows = [dict(r) for r in conn.execute("SELECT * FROM surahs ORDER BY surah_id")]
        by_surah = defaultdict(list)
        v_by_surah = defaultdict(list)
        for token in reference["tokens"]:
            by_surah[token["surah_id"]].append(token)
        for verse in reference["verses"]:
            v_by_surah[verse["surah_id"]].append(verse)
        expected = [{"surah_id": s, "ayah_count": len(v_by_surah[s]), "token_count": len(ts),
                     "numbered_token_count": sum(not t["is_opening_basmala"] for t in ts),
                     "first_verse_id": v_by_surah[s][0]["verse_id"], "last_verse_id": v_by_surah[s][-1]["verse_id"]}
                    for s, ts in sorted(by_surah.items())]
        _record(checks, "aggregates.surahs", rows == expected, "Все агрегаты сур сверены с независимыми исходными токенами и аятами.", surahs=len(rows))
    for scope in ("file", "numbered"):
        selected = [t for t in reference["tokens"] if scope == "file" or not t["is_opening_basmala"]]
        if "token_scopes" in tables:
            rows = [dict(r) for r in conn.execute("SELECT * FROM token_scopes WHERE scope=? ORDER BY scope_token_index", (scope,))]
            verses, surahs = Counter(), Counter()
            errors = []
            for index, (token, row) in enumerate(zip(selected, rows), 1):
                verses[token["verse_id"]] += 1
                surahs[token["surah_id"]] += 1
                expected = {"token_id": token["token_id"], "scope": scope, "scope_token_index": index,
                            "scope_token_in_ayah": verses[token["verse_id"]], "scope_token_in_surah": surahs[token["surah_id"]]}
                if row != expected and len(errors) < 20:
                    errors.append(token["token_id"])
            _record(checks, f"positions.{scope}", len(rows) == len(selected) and not errors,
                    "Плотные индексы корпуса, аята и суры независимо пересчитаны после исключения префиксов.", mismatches=errors)
        for profile in PROFILES:
            if "vocabulary" not in tables:
                continue
            groups = defaultdict(list)
            for index, token in enumerate(selected, 1):
                groups[token[profile]].append((index, token))
            rows = [dict(r) for r in conn.execute("SELECT * FROM vocabulary WHERE profile_id=? AND scope=? ORDER BY rank", (profile, scope))]
            cumulative = 0
            errors = []
            for rank, ((form, members), row) in enumerate(zip(sorted(groups.items(), key=lambda p: (-len(p[1]), p[0])), rows), 1):
                cumulative += len(members)
                expected = {"profile_id": profile, "scope": scope, "form": form, "frequency": len(members), "rank": rank,
                            "ayah_count": len({t["verse_id"] for _, t in members}), "surah_count": len({t["surah_id"] for _, t in members}),
                            "first_token_id": members[0][1]["token_id"], "last_token_id": members[-1][1]["token_id"],
                            "first_scope_token": members[0][0], "last_scope_token": members[-1][0]}
                real_expected = {"share": len(members) / len(selected), "cumulative_share": cumulative / len(selected),
                                 "mean_normalized_position": sum(i - 1 for i, _ in members) / len(members) / max(1, len(selected) - 1)}
                good = all(row.get(k) == v for k, v in expected.items())
                good &= all(isinstance(row.get(k), (float, int)) and math.isclose(row[k], v, rel_tol=1e-10, abs_tol=1e-12) for k, v in real_expected.items())
                if not good and len(errors) < 20:
                    errors.append(form)
            _record(checks, f"vocabulary.{scope}.{profile}", len(rows) == len(groups) and not errors,
                    "Каждая словарная форма, частота, ранг, охват, крайние адреса и нормированная средняя позиция сверены с исходником.", mismatches=errors, forms=len(rows))
            if {"occurrences", "token_scopes"} <= tables:
                bad = conn.execute(f"""SELECT COUNT(*) FROM occurrences o
                    JOIN vocabulary v USING(unit_id) LEFT JOIN tokens t USING(token_id)
                    LEFT JOIN token_scopes s ON s.token_id=o.token_id AND s.scope=v.scope
                    WHERE v.profile_id=? AND v.scope=? AND
                    (t.token_id IS NULL OR s.token_id IS NULL OR t.{profile}!=v.form OR o.scope_token_index!=s.scope_token_index)""", (profile, scope)).fetchone()[0]
                totals = dict(conn.execute("SELECT v.unit_id, COUNT(o.token_id) FROM vocabulary v LEFT JOIN occurrences o USING(unit_id) WHERE v.profile_id=? AND v.scope=? GROUP BY v.unit_id", (profile, scope)))
                counts_good = all(totals.get(row["unit_id"]) == row["frequency"] for row in rows)
                actual_n = sum(totals.values())
                _record(checks, f"occurrences.{scope}.{profile}", not bad and counts_good and actual_n == len(selected),
                        "Полный индекс вхождений: форма каждого токена, область, позиция, частота каждой единицы и общий знаменатель.", occurrences=actual_n, bad_links=bad)


def _validate_basmala(reference: dict, checks: list):
    by_id = {v["verse_id"]: v for v in reference["verses"]}
    tokens = reference["tokens"]
    by_verse = defaultdict(list)
    for t in tokens:
        by_verse[t["verse_id"]].append(t)
    phrase = ["بسم", "الله", "الرحمن", "الرحيم"]
    special = {
        "1:1": "Нумерованная басмала 1:1 сохраняется в numbered.",
        "9:1": "В начале 9:1 нет басмалы и нет исключённых токенов.",
        "27:30": "Внутренняя басмала 27:30 сохраняется целиком.",
    }
    for verse_id, method in special.items():
        rows = by_verse.get(verse_id, [])
        forms = [r["plain"] for r in rows]
        retained = all(not r["is_opening_basmala"] for r in rows)
        contains = any(forms[i:i + 4] == phrase for i in range(max(0, len(forms) - 3)))
        _record(checks, f"basmala.{verse_id}", bool(rows) and retained and (not contains if verse_id == "9:1" else contains), method)
    flags = [t for t in tokens if t["is_opening_basmala"]]
    opening_ids = sorted({t["verse_id"] for t in flags})
    eligible = [v for v in reference["verses"] if v["ayah_id"] == 1 and v["surah_id"] not in (1, 9)]
    expected_openings = {f"{s}:1" for s in range(2, 115) if s != 9}
    _record(checks, "basmala.prefixes", len(flags) == 448 and set(opening_ids) == expected_openings, "В данном файле 112 начальных префиксов по четыре токена; raw-варианты 95:1 и 97:1 не теряются.", excluded_tokens=len(flags), prefix_verses=len(opening_ids), expected_prefixes=112)
    modified = []
    reference_raw = [t["raw"] for t in by_verse.get("1:1", [])]
    for verse_id in opening_ids:
        rows = by_verse[verse_id][:4]
        if [t["raw"] for t in rows] != reference_raw:
            modified.append(verse_id)
    variants_valid = set(modified) == {"95:1", "97:1"}
    for verse_id in modified:
        raw_forms = [t["raw"] for t in by_verse[verse_id][:4]]
        variants_valid &= raw_forms[0].count("\u0651") == reference_raw[0].count("\u0651") + 1 and raw_forms[0].replace("\u0651", "") == reference_raw[0] and raw_forms[1:] == reference_raw[1:]
    _record(checks, "basmala.orthographic_variants", variants_valid, "Только 95:1 и 97:1 отличаются от 1:1 дополнительной шаддой первого токена; исходник не исправляется.", raw_variant_verses=modified)


def _validate_unicode(checks: list):
    samples = {"ا\u0654": "أ", "ا\u0655": "إ", "ا\u0653": "آ", "ب\u0654": "ب\u0654", "ب\u0655": "ب\u0655", "ب\u0653": "ب\u0653", "ة": "ة", "ى": "ى", "بِّ": "ب", "ـ": ""}
    failures = [{"raw": raw, "expected": plain, "observed": independent_profiles(raw)["plain"]} for raw, plain in samples.items() if independent_profiles(raw)["plain"] != plain]
    _record(checks, "unicode.hamza_madda_preserved", not failures, "Контроль спецификации: NFC может составить хамзу с алифом; U+0653..U+0655 не удаляются, ة и ى в plain различаются.", samples=len(samples), failures=failures, limitation_ru="Эти синтетические примеры проверяют независимый оракул; производственная нормализация дополнительно покрывается tests/test_core.py.")


def _surface_totals(tokens, profile):
    """Whole-string codepoint counts; no production metric implementation."""
    text = "".join(t[profile] for t in tokens)
    categories = Counter(ud.category(c) for c in text)
    return {"tokens": len(tokens), "types": len({t[profile] for t in tokens}), "codepoints": len(text),
            "letters": categories["Lo"], "marks": sum(v for k, v in categories.items() if k.startswith("M")),
            "harakat": sum("\u064b" <= c <= "\u0652" for c in text), "tatweel": text.count("ـ"),
            "superscript_alef": text.count("\u0670")}


def _validate_core_exports(root, reference, source_bytes, checks):
    expected_totals = {}
    by_verse = defaultdict(list)
    for token in reference["tokens"]:
        by_verse[token["verse_id"]].append(token)
    for scope in ("file", "numbered"):
        selected = [t for t in reference["tokens"] if scope == "file" or not t["is_opening_basmala"]]
        for profile in PROFILES:
            expected_totals[scope, profile] = _surface_totals(selected, profile)
            p = root / f"results/tables/core_vocabulary_{profile}_{scope}.csv.gz"
            if p.exists():
                rows = list(_rows(p))
                actual = {r["form"]: int(r["frequency"]) for r in rows}
                expected = dict(Counter(t[profile] for t in selected))
                _record(checks, f"csv.vocabulary.{scope}.{profile}", actual == expected and len(rows) == len(actual),
                        "Все частоты CSV.gz сверены с независимыми исходными вхождениями.", rows=len(rows))
            else:
                _record(checks, f"csv.vocabulary.{scope}.{profile}", None, "Частотный CSV.gz ещё не создан.")
    for filename, kind in (("core_verse_lengths.csv.gz", "verse"), ("core_surah_lengths.csv", "surah")):
        path = root / "results/tables" / filename
        if not path.exists():
            _record(checks, f"csv.lengths.{kind}", None, "Таблица длин ещё не создана.")
            continue
        rows = list(_rows(path))
        expected_rows = {}
        for scope in ("file", "numbered"):
            groups = defaultdict(list)
            for t in reference["tokens"]:
                if scope == "file" or not t["is_opening_basmala"]:
                    groups[t["verse_id"] if kind == "verse" else str(t["surah_id"])].append(t)
            for key, ts in groups.items():
                for profile in PROFILES:
                    expected_rows[scope, profile, key] = _surface_totals(ts, profile)
        errors, seen = [], set()
        for row in rows:
            key = row["scope"], row["profile_id"], row["verse_id" if kind == "verse" else "surah_id"]
            expected = expected_rows.get(key)
            good = expected is not None and key not in seen
            if expected:
                good &= all(int(row[field]) == value for field, value in expected.items() if field != "types")
                good &= int(row["unique_forms"]) == expected["types"]
            if not good and len(errors) < 20:
                errors.append(key)
            seen.add(key)
        _record(checks, f"csv.lengths.{kind}", not errors and seen == set(expected_rows),
                "Длины каждой структурной единицы по 4 профилям × 2 областям сверены с raw-токенами; графемы отдельно этим оракулом не проверяются.", rows=len(rows), mismatches=errors)
    path = root / "results/tables/core_basmala_sensitivity.csv"
    if path.exists():
        rows = list(_rows(path))
        errors = [f"{r['scope']}.{r['profile_id']}" for r in rows if any(int(r[k]) != v for k, v in expected_totals[r["scope"], r["profile_id"]].items() if k != "types")]
        _record(checks, "csv.scope_totals", len(rows) == 8 and not errors, "Все профильные суммы токенов, букв, знаков, кодовых точек и явно удаляемых знаков.", mismatches=errors)
    else:
        _record(checks, "csv.scope_totals", None, "Сводная таблица областей ещё не создана.")
    path = root / "results/tables/core_source_offsets_bytes.csv.gz"
    if path.exists():
        token_lookup = {t["token_id"]: t for t in reference["tokens"]}
        wanted = {t[k] for t in reference["tokens"] for k in ("raw_start", "raw_end")}
        byte_positions, byte_offset = {}, 0
        for cp, char in enumerate(source_bytes.decode("utf-8")):
            if cp in wanted:
                byte_positions[cp] = byte_offset
            byte_offset += len(char.encode("utf-8"))
        byte_positions[len(source_bytes.decode("utf-8"))] = byte_offset
        errors, seen = [], set()
        for row in _rows(path):
            token = token_lookup.get(row["token_id"])
            good = token is not None and row["token_id"] not in seen
            if token:
                good &= int(row["codepoint_start"]) == token["raw_start"] and int(row["codepoint_end"]) == token["raw_end"]
                good &= int(row["byte_start"]) == byte_positions[token["raw_start"]] and int(row["byte_end"]) == byte_positions[token["raw_end"]]
                good &= source_bytes[int(row["byte_start"]):int(row["byte_end"])] == token["raw"].encode("utf-8")
            if not good and len(errors) < 20:
                errors.append(row["token_id"])
            seen.add(row["token_id"])
        _record(checks, "source.byte_offsets", not errors and seen == set(token_lookup), "Байтовые координаты каждого токена проверяются непосредственно на неизменных исходных байтах.", mismatches=errors)
    else:
        _record(checks, "source.byte_offsets", None, "Таблица байтовых смещений ещё не создана.")
    return expected_totals


def _validate_statistics(root, config, checks):
    path = root / "results/tables/statistical_tests.csv"
    null_path = root / "results/tables/statistical_null_draws.csv.gz"
    if not path.exists() or not null_path.exists():
        _record(checks, "statistics.saved_draws", None, "Тесты и полный набор нулевых симуляций ещё не созданы.")
        return
    tests = list(_rows(path))
    draws = defaultdict(dict)
    duplicate_draws = []
    for row in _rows(null_path):
        hid, i = row["hypothesis_id"], int(row["simulation"])
        if i in draws[hid]:
            duplicate_draws.append((hid, i))
        draws[hid][i] = float(row["value"])
    pvalues, errors = [], []
    for test in tests:
        values = draws.get(test["hypothesis_id"], {})
        n = int(test["simulations"])
        observed = float(test["observed"])
        if set(values) != set(range(1, n + 1)):
            errors.append(test["hypothesis_id"] + ":draw_indices")
            pvalues.append(1.0)
            continue
        arr = list(values.values())
        varies = max(arr) - min(arr) > 1e-12
        extreme = sum(v <= observed if test["tail"] == "less" else v >= observed for v in arr)
        p = (extreme + 1) / (n + 1) if varies else 1.0
        pvalues.append(p)
        mean = math.fsum(arr) / n
        if not math.isclose(p, float(test["p_value"]), abs_tol=1e-12) or not math.isclose(mean, float(test["null_mean"]), abs_tol=1e-9) or not math.isclose(observed - mean, float(test["effect"]), abs_tol=1e-9):
            errors.append(test["hypothesis_id"] + ":p_mean_effect")
    # Scalar Holm recurrence, independent of the production NumPy algorithm.
    adjusted = [0.0] * len(tests)
    running = 0.0
    for i, k in enumerate(sorted(range(len(tests)), key=lambda k: pvalues[k])):
        running = max(running, (len(tests) - i) * pvalues[k])
        adjusted[k] = min(1.0, running)
    errors += [t["hypothesis_id"] + ":holm" for t, p in zip(tests, adjusted) if not math.isclose(float(t["p_adjusted"]), p, abs_tol=1e-12)]
    _record(checks, "statistics.saved_draws", bool(tests) and not errors and not duplicate_draws and set(draws) == {t["hypothesis_id"] for t in tests},
            "По всем сохранённым симуляциям независимо пересчитаны p=(b+1)/(B+1), включение равенств, среднее, эффект и единая скалярная поправка Holm.", tests=len(tests), simulations=sum(map(len, draws.values())), mismatches=errors[:20], duplicate_draws=duplicate_draws[:20])


def _validate_catalogues(root: Path, checks: list):
    coverage_files = sorted(p for p in (root / "results").glob("*_coverage.json") if not p.name.startswith("validation"))
    rows = []
    for p in coverage_files:
        obj = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(obj, list):
            rows.extend({"file": str(p.relative_to(root)), **r} for r in obj)
        elif isinstance(obj, dict):
            rows.extend({"file": str(p.relative_to(root)), **r} for r in obj.get("coverage", []))
    directions = {r.get("direction") for r in rows}
    _record(checks, "coverage.directions_A_J", set("ABCDEFGHIJ") <= directions if coverage_files else None, "Для каждого направления A–J существует строка результата либо конкретного ограничения.", directions=sorted(str(d) for d in directions), files=[str(p.relative_to(root)) for p in coverage_files])
    missing = [r for r in rows if not all(k in r for k in ("analysis", "status", "universe", "tested", "limitations"))]
    _record(checks, "coverage.explicit_bounds", not missing if rows else None, "У каждой записи указаны статус, пространство поиска, фактическое покрытие и ограничения.", missing_fields=missing[:10])
    hypotheses = []
    for p in sorted((root / "results/hypotheses").glob("*.json")):
        if p.name == "registry.json":
            continue  # This is a union of module files, not a new hypothesis source.
        obj = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(obj, list):
            hypotheses.extend({"file": str(p.relative_to(root)), **r} for r in obj)
    absent_evidence = []
    for h in hypotheses:
        path = h.get("evidence_path")
        if isinstance(path, str) and path and not (root / path.split("#")[0]).exists():
            absent_evidence.append({"hypothesis_id": h.get("hypothesis_id"), "path": path})
    ids = [h.get("hypothesis_id") for h in hypotheses]
    _record(checks, "hypotheses.unique_ids", len(ids) == len(set(ids)) if ids else None, "Идентификаторы гипотез уникальны во всех модулях.", hypotheses=len(ids))
    _record(checks, "hypotheses.evidence_exists", not absent_evidence if hypotheses else None, "Пути evidence_path разрешаются в реальные файлы результатов.", missing=absent_evidence)
    return rows, hypotheses


def _same_export_value(actual, expected):
    if expected in (None, ""):
        return actual in (None, "")
    if isinstance(actual, (int, float)) and not isinstance(actual, bool):
        try:
            return math.isclose(actual, float(expected), rel_tol=1e-12, abs_tol=1e-12)
        except (ValueError, TypeError):
            return False
    return actual == str(expected)


def _validate_exports(root: Path, reference: dict, totals: dict, checks: list):
    """Reconcile serialized deliverables with their declared source tables."""
    workbooks = sorted((root / "exports").glob("*.xlsx"))
    if not workbooks:
        _record(checks, "export.xlsx_readable", None, "Книга Excel ещё не создана.")
    for path in workbooks:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
        except (ImportError, OSError, ValueError) as exc:
            _record(checks, "export.xlsx_readable", False, "Не удалось прочитать книгу через openpyxl.", error=str(exc))
            continue
        numeric, textual_addresses = 0, 0
        bad_addresses, bad_layout = [], []
        sheet_info = []
        for sheet in wb:
            sheet_info.append({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column})
            if sheet.freeze_panes != "A2" or not sheet.auto_filter.ref:
                bad_layout.append(sheet.title)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        numeric += 1
                    if isinstance(cell.value, str) and re.fullmatch(r"\d{1,3}:\d{1,3}(?::\d{1,3})?", cell.value):
                        textual_addresses += 1
                        if cell.data_type != "s" or cell.number_format != "@":
                            bad_addresses.append({"sheet": sheet.title, "cell": cell.coordinate})
        _record(checks, "export.xlsx_readable", bool(wb.sheetnames) and numeric > 0 and not bad_addresses,
                "XLSX читается; найденные числовые значения — числа, адреса — явный текст с форматом @.", path=str(path.relative_to(root)), sheets=sheet_info, numeric_cells=numeric, address_text_cells=textual_addresses, bad_addresses=bad_addresses[:20])
        _record(checks, "export.xlsx_filters", not bad_layout, "На всех листах закреплён заголовок и установлен фильтр.", bad_sheets=bad_layout)
        reconciled, errors, oversized = [], [], []
        if "Указатель листов" in wb.sheetnames:
            index = list(wb["Указатель листов"].values)
            for source in (dict(zip(index[0], row)) for row in index[1:]):
                relative = source.get("Источник", "")
                csv_path = root / relative
                if not relative.endswith(".csv") or not csv_path.exists():
                    continue
                title = source["Лист"]
                actual_rows = list(wb[title].values)
                header, actual_rows = actual_rows[0], actual_rows[1:]
                expected_rows = list(_rows(csv_path))
                good = len(actual_rows) == len(expected_rows) == int(source["Строк данных"])
                for observed, expected in zip(actual_rows, expected_rows):
                    for k,v in zip(header,observed):
                        value=expected.get(k)
                        if isinstance(value,str) and len(value)>32767:
                            notice=f"Полное значение доступно в {relative}; длина {len(value)} символов, превышает лимит ячейки XLSX."
                            good &= v==notice
                            oversized.append({'source':relative,'column':k,'characters':len(value)})
                        else:
                            good &= _same_export_value(v,value)
                if good:
                    reconciled.append(relative)
                else:
                    errors.append({"sheet": title, "source": relative, "observed_rows": len(actual_rows), "expected_rows": len(expected_rows)})
        _record(checks, "export.xlsx_csv_consistency", bool(reconciled) and not errors,
                "Все листы с объявленным CSV-источником сопоставлены ячейка за ячейкой; адреса не преобразуются в числа. Для ячеек свыше лимита32767 проверено точное уведомление со ссылкой на полный исходник.", reconciled_tables=len(reconciled), oversized_cells_with_explicit_source=oversized, mismatches=errors[:20])
        wb.close()
    metrics_path = root / "results/tables/presentation_corpus_metrics.csv"
    if metrics_path.exists():
        rows = list(_rows(metrics_path))
        errors, seen = [], set()
        for row in rows:
            key = row["scope"], row["profile_id"]
            expected = totals.get(key)
            good = expected is not None and key not in seen
            if expected:
                good &= all(int(row[k]) == expected[k] for k in ("tokens", "types", "letters", "marks"))
                good &= int(row["verses"]) == len(reference["verses"]) and int(row["surahs"]) == len({v["surah_id"] for v in reference["verses"]})
            if not good:
                errors.append(key)
            seen.add(key)
        _record(checks, "export.metrics_source", not errors and seen == set(totals), "Сводные метрики отчёта и XLSX по всем профилям/областям сверены с независимым чтением исходника.", mismatches=errors)
    else:
        _record(checks, "export.metrics_source", None, "Сводные метрики представления ещё не созданы.")
    dashboard = root / "dashboard/index.html"
    if not dashboard.exists():
        _record(checks, "export.html_source", None, "Переносимый HTML ещё не создан.")
    else:
        html = dashboard.read_text(encoding="utf-8")
        match = re.search(r'<script[^>]*id=["\']data-json["\'][^>]*>(.*?)</script>', html, re.S)
        data = json.loads(match.group(1)) if match else None
        _record(checks, "export.html_payload", data is not None, "JSON-набор HTML извлекается из script#data-json.")
        if data is not None:
            errors = []
            for i, (row, token) in enumerate(zip(data["tokens"], reference["tokens"])):
                expected = [token["global_ayah"] - 1, token["token_in_ayah"], token["global_token"], token["token_in_surah"], token["raw_start"], token["raw_end"]]
                good = row[:6] == expected and row[10:] == [token["is_opening_basmala"], token["letter_count"], token["mark_count"]]
                good &= all(data["forms"][profile][row[6 + p]] == token[profile] for p, profile in enumerate(PROFILES))
                if not good and len(errors) < 20:
                    errors.append(token["token_id"])
            fields = ("verse_id", "surah_id", "ayah_id", "global_ayah", "raw_text", "raw_start", "raw_end", "token_count", "letter_count", "mark_count_all")
            verses_good = data["verses"] == [[v[k] for k in fields] for v in reference["verses"]]
            _record(checks, "export.html_source", len(data["tokens"]) == len(reference["tokens"]) and not errors and verses_good,
                    "Каждый встроенный токен, 4 поверхности, адрес, басмальный флаг и полный аят сверены непосредственно с независимым исходником.", tokens=len(data["tokens"]), verses=len(data["verses"]), mismatches=errors)
            expected_metrics = list(_rows(metrics_path)) if metrics_path.exists() else []
            metrics_good = len(data["metrics"]) == len(expected_metrics) and all(all(_same_export_value(a[k], b[k]) for k in b) for a, b in zip(data["metrics"], expected_metrics))
            _record(checks, "export.html_metrics_csv", metrics_good, "Встроенные сводные метрики HTML совпадают с CSV, проверенным по исходнику.")
            figure_errors = []
            for figure in data.get("figures", []):
                for key in ("png", "svg", "data"):
                    if not (root / figure[key]).is_file():
                        figure_errors.append(figure[key])
            _record(checks, "export.figure_sources", bool(data.get("figures")) and not figure_errors, "Для каждой опубликованной статической визуализации существуют PNG, SVG и CSV исходных чисел.", missing=figure_errors)
    report = root / "report/REPORT.md"
    if report.exists():
        text = report.read_text(encoding="utf-8")
        passport={parts[0].strip('`'):parts[1:] for line in text.splitlines() if line.startswith('|') and len(parts:=[p.strip() for p in line.strip('|').split('|')])==5}
        good = all(passport.get(profile)==[str(totals['file',profile]['tokens']),str(totals['file',profile]['types']),str(totals['numbered',profile]['tokens']),str(totals['numbered',profile]['types'])] for profile in PROFILES)
        _record(checks, "export.report_metrics", good, "Все четыре строки профильного паспорта текстового отчёта сверены с независимыми частотами.")
    else:
        _record(checks, "export.report_metrics", None, "Текстовый отчёт ещё не создан.")


def _validate_discoveries(root, reference, checks):
    """Recompute finite dictionaries from the independently parsed original."""
    byverse=defaultdict(list)
    token_map={t['token_id']:t for t in reference['tokens']}
    for t in reference['tokens']:byverse[t['verse_id']].append(t)
    path=root/'results/tables/exploration_word_ngrams.csv.gz'
    if path.exists():
        rows=list(_rows(path));ns=sorted({int(r['n']) for r in rows})
        counts=Counter()
        for vv in byverse.values():
            words=[t['plain'] for t in vv]
            for n in ns:
                # Simple zip of shifted copies, separate from production index.
                counts.update((n,' '.join(x)) for x in zip(*(words[i:] for i in range(n))))
        actual={(int(r['n']),r['sequence']):int(r['frequency']) for r in rows}
        _record(checks,'discoveries.full_word_ngrams',actual==dict(counts) and len(rows)==len(actual),
                'Полный словарь всех заявленных словесных n-грамм заново получен из независимых исходных токенов с границами аятов.',types=len(actual),occurrences=sum(actual.values()))
        ids={r['ngram_id']:(int(r['n']),r['sequence']) for r in rows}
        observed=Counter();errors=[]
        for r in _rows(root/'results/tables/exploration_word_ngram_occurrences.csv.gz'):
            start=token_map.get(r['start_token_id']);end=token_map.get(r['end_token_id']);key=ids.get(r['ngram_id'])
            good=bool(start and end and key)
            if good:
                vv=byverse[start['verse_id']];i=start['token_in_ayah']-1
                part=vv[i:i+key[0]]
                good=(len(part)==key[0] and part[-1]['token_id']==end['token_id'] and
                      ' '.join(t['plain'] for t in part)==key[1] and start['verse_id']==r['verse_id'] and
                      int(r['global_token'])==start['global_token'])
            if not good and len(errors)<20:errors.append(r)
            observed[r['ngram_id']]+=1
        _record(checks,'discoveries.ngram_all_coordinates',not errors and all(observed[r['ngram_id']]==int(r['frequency']) for r in rows) and set(observed)==set(ids),
                'Каждое начало/окончание полного списка n-грамм восстановлено по исходным token_id; агрегаты вхождений сверены со словарём.',mismatches=errors)
    else:_record(checks,'discoveries.full_word_ngrams',None,'Поиск повторов ещё не рассчитан.')
    path=root/'results/tables/exploration_exact_verse_repeats.csv'
    if path.exists():
        expected={}
        for profile in PROFILES:
            for scope in ['file','numbered']:
                groups=defaultdict(list)
                for v in reference['verses']:
                    vv=[t for t in byverse[v['verse_id']] if scope=='file' or not t['is_opening_basmala']]
                    key=v['raw_text'] if profile=='raw' and scope=='file' else ' '.join(t[profile] for t in vv)
                    groups[key].append(v['verse_id'])
                expected.update({(profile,scope,key):ids for key,ids in groups.items() if len(ids)>1})
        actual={(r['profile_id'],r['scope'],r['sequence']):r['verse_ids'].split('|') for r in _rows(path)}
        _record(checks,'discoveries.full_verse_repeats',actual==expected,'Полный реестр повторённых аятов во всех 8 вариантах независимо пересчитан вместе со всеми адресами.',groups=len(actual))
    path=root/'results/tables/exploration_long_repeats.csv'
    if path.exists():
        errors=[];rows=list(_rows(path))
        for r in rows:
            starts=r['start_token_ids'].split('|');ends=r['end_token_ids'].split('|');n=int(r['token_length'])
            good=len(starts)==len(ends)==int(r['frequency'])
            for start,end in zip(starts,ends):
                t=token_map[start];part=byverse[t['verse_id']][t['token_in_ayah']-1:t['token_in_ayah']-1+n]
                good &= len(part)==n and part[-1]['token_id']==end and ' '.join(x['plain'] for x in part)==r['sequence']
            if not good:errors.append(r['sequence'])
        _record(checks,'discoveries.long_repeat_source_contexts',bool(rows) and not errors,'Каждый найденный длинный повтор восстановлен по исходным адресам; это проверка кандидатов, не доказательство полноты эвристического поиска.',groups=len(rows),mismatches=errors[:20])


def _validate_morphology(root, reference, checks):
    path=root/'results/tables/morphology_segments.csv.gz'
    if not path.exists():
        _record(checks,'morphology.join_conservation',None,'Морфологические таблицы ещё не рассчитаны.');return
    original=set()
    for line in (root/'data/external/quranic-corpus-morphology-0.4.txt').read_text(encoding='utf-8').splitlines():
        if line.startswith('('):original.add(line.split('\t',1)[0].strip('()'))
    rows=list(_rows(path));ids=[r['segment_id'] for r in rows]
    token_map={t['token_id']:t for t in reference['tokens']};errors=[]
    for r in rows:
        for tid in r['token_ids'].split('|'):
            t=token_map.get(tid)
            if not t or t['is_opening_basmala'] or t['verse_id']!=r['verse_id']:
                if len(errors)<20:errors.append(tid)
    _record(checks,'morphology.join_conservation',not errors and len(ids)==len(set(ids)) and set(ids)<=original,
            'Все перенесённые segment_id уникальны и присутствуют в неизменённом QAC; связанные токены существуют, в том же аяте и не являются вступительными префиксами.',segments=len(rows),mismatches=errors)
    errors=[]
    for field,filename in [('lemma','lemma'),('root','root'),('pos','pos'),('segment_arabic','segment')]:
        counter=Counter(r[field] for r in rows if r[field]);links=defaultdict(set)
        for r in rows:
            if r[field]:links[r[field]].update(r['token_ids'].split('|'))
        got=list(_rows(root/f'results/tables/morphology_{filename}_frequencies.csv'))
        if {r['unit']:int(r['segment_count']) for r in got}!=dict(counter):errors.append(field+':segments')
        if any(int(r['token_count'])!=len(links[r['unit']]) for r in got):errors.append(field+':linked_tokens')
    _record(checks,'morphology.frequency_conservation',not errors,'Частоты всех лемм, корней, POS и форм сегментов независимо агрегированы; сегменты и связанные токены имеют отдельные счётчики.',mismatches=errors)


def run(root: Path, config: dict) -> dict:
    root = Path(root)
    checks = []
    source_path = root / config.get("input", "Quran_text.txt")
    source_bytes = source_path.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    expected_hash = config.get("validation", {}).get("source_sha256", SOURCE_SHA256)
    _record(checks, "source.sha256", source_hash == expected_hash, "SHA-256 сравнивается с хешем исходного файла, независимо зафиксированным перед обработкой.", observed=source_hash, expected=expected_hash)
    raw_path = root / "data/raw" / source_path.name
    _record(checks, "source.immutable_copy", raw_path.exists() and raw_path.read_bytes() == source_bytes, "Побайтовая идентичность оригинала и неизменной копии.", path=str(raw_path.relative_to(root)))
    text = source_bytes.decode("utf-8")
    reference = independent_source(text)
    _record(checks, "source.independent_parse", not reference["issues"], "Независимый разбор каждой строки, включая числовые метаданные и служебные знаки.", issues=reference["issues"][:20], lines=len(text.splitlines()), verses=len(reference["verses"]))
    actual_ids = [(v["surah_id"], v["ayah_id"]) for v in reference["verses"]]
    expected_ids = []
    for s in sorted({s for s, _ in actual_ids}):
        expected_ids.extend((s, a) for a in range(1, max(a for ss, a in actual_ids if ss == s) + 1))
    _record(checks, "source.record_order", actual_ids == expected_ids and {s for s, _ in actual_ids} == set(range(1, 115)), "По данным файла: 114 последовательных сур, уникальные и последовательные номера аятов без пропусков; полнота внешней редакции этим не доказывается.")
    _validate_unicode(checks)
    _validate_basmala(reference, checks)
    db = validate_database(root / "data/processed/corpus.sqlite", text, reference, checks)
    expected_totals = _validate_core_exports(root, reference, source_bytes, checks)
    _validate_statistics(root, config, checks)
    _validate_discoveries(root, reference, checks)
    _validate_morphology(root, reference, checks)
    coverage, hypotheses = _validate_catalogues(root, checks)
    _validate_exports(root, reference, expected_totals, checks)
    independent_counts = {
        "sha256": source_hash, "bytes": len(source_bytes), "unicode_codepoints_file": len(text),
        "surahs": len({v["surah_id"] for v in reference["verses"]}), "verses": len(reference["verses"]),
        "tokens_file": len(reference["tokens"]), "tokens_numbered": sum(not t["is_opening_basmala"] for t in reference["tokens"]),
        "base_letters_file": sum(t["letter_count"] for t in reference["tokens"]),
        "base_letters_numbered": sum(t["letter_count"] for t in reference["tokens"] if not t["is_opening_basmala"]),
        "marks_tokens_file": sum(t["mark_count"] for t in reference["tokens"]),
        "marks_verses_file": sum(v["mark_count_all"] for v in reference["verses"]),
        "profile_counts": db.get("profile_counts", {}),
    }
    statuses = Counter(c["status"] for c in checks)
    status = "fail" if statuses["fail"] else "incomplete" if statuses["not_run"] else "pass"
    summary = {"status": status, "checks": len(checks), "statuses": dict(statuses), "independent_counts": independent_counts, "failed_checks": [c["check_id"] for c in checks if c["status"] == "fail"], "not_run_checks": [c["check_id"] for c in checks if c["status"] == "not_run"], "checks_path": "results/validation_checks.json", "method_ru": "Независимое чтение raw-строк, разбиение по пробелам с фильтром букв Lo и отдельные SQL/экспортные сверки; core не импортируется.", "limitations_ru": ["Проверка не подтверждает происхождение, чтение и внешнюю редакционную полноту файла.", "Визуальное отображение HTML проверяется отдельно; openpyxl проверяет структуру и типы XLSX, но не вид книги в Excel.", "Проверка статусов A–J подтверждает наличие явной карты, но не доказывает семантическую полноту каждого направления."]}
    (root / "results").mkdir(exist_ok=True)
    for name, obj in (("validation_summary.json", summary), ("validation_checks.json", checks), ("validation_independent_counts.json", independent_counts)):
        (root / "results" / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary
